#!/usr/bin/env python3
"""Fail when the final Word report diverges from the reference bidi pattern."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


HEBREW = re.compile(r"[\u0590-\u05ff]")
LATIN = re.compile(r"[A-Za-z]")
FORBIDDEN = {"\u200e", "\u200f", "\u202a", "\u202b", "\u202c", "\u2066", "\u2067", "\u2068", "\u2069"}
DUPLICATE_TRANSLATIONS = (
    "הפרש קוד (code diff)",
    "הפצה לאחור",
    "אנטרופיה צולבת",
    "פירוק לאסימונים",
    "מקטעים בני כמה אסימונים",
    "הטמעות",
    "מנגנון קשב",
    "הטמעת מיקום",
    "נשירת יחידות",
    "דעיכת משקלים",
    "תיוג חלש",
    "התיוג החלש",
    "מקטעי תווים קצרים",
    "הטמעה (embedding)",
    "דעיכת המשקלים",
    "קשב עצמי",
    "שיעור הדיוק (accuracy)",
    "דיוק ההתראות (precision)",
    "שיעור הכיסוי (recall)",
)


def paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def main() -> None:
    path = Path(sys.argv[1] if len(sys.argv) > 1 else "report/DiffIntent_RTL_Final_Report_Hebrew.docx")
    errors = []
    for index, paragraph in enumerate(paragraphs(Document(path))):
        text = paragraph.text
        if any(mark in text for mark in FORBIDDEN):
            errors.append(f"paragraph {index}: hidden direction mark")
        if re.search(r"[\u0590-\u05ff][A-Za-z]|[A-Za-z][\u0590-\u05ff]", text):
            errors.append(f"paragraph {index}: Hebrew and Latin letters touch without a separator")
        for duplicate in DUPLICATE_TRANSLATIONS:
            if duplicate in text:
                errors.append(f"paragraph {index}: duplicate Hebrew/English term: {duplicate}")
        if HEBREW.search(text):
            p_pr = paragraph._p.get_or_add_pPr()
            if p_pr.find(qn("w:bidi")) is None:
                errors.append(f"paragraph {index}: Hebrew paragraph lacks w:bidi")
            if paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
                errors.append(f"paragraph {index}: RTL paragraph uses physical RIGHT alignment")
        for run_index, run in enumerate(paragraph.runs):
            run_text = run.text
            r_pr = run._r.get_or_add_rPr()
            rtl = r_pr.find(qn("w:rtl"))
            if LATIN.search(run_text) and not HEBREW.search(run_text):
                if rtl is not None:
                    errors.append(f"paragraph {index}, run {run_index}: Latin run has w:rtl")
                lang = r_pr.find(qn("w:lang"))
                if lang is None or lang.get(qn("w:val")) != "en-US":
                    errors.append(f"paragraph {index}, run {run_index}: Latin run lacks en-US language")
            elif HEBREW.search(run_text):
                if rtl is None or rtl.get(qn("w:val"), "1") in {"0", "false", "off"}:
                    errors.append(f"paragraph {index}, run {run_index}: Hebrew run lacks active w:rtl")
    assert not errors, "\n" + "\n".join(errors[:40])
    print(f"PASS: {path}")


if __name__ == "__main__":
    main()
