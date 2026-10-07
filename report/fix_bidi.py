#!/usr/bin/env python3
"""Normalize the final report to Word's native Hebrew/English bidi model."""

import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


HEBREW = re.compile(r"[\u0590-\u05ff]")
LATIN = re.compile(r"[A-Za-z]")
ENGLISH_EXPRESSION = re.compile(
    r"[A-Za-z][A-Za-z0-9_]*(?:[-./+][A-Za-z0-9_]+)*(?:[ \u00a0][A-Za-z][A-Za-z0-9_]*(?:[-./+][A-Za-z0-9_]+)*)*"
)
DELETE_MARKS = str.maketrans("", "", "\u200e\u200f\u202a\u202b\u202c\u2066\u2067\u2068\u2069")
REPLACEMENTS = (
    ("הפרש קוד (code diff)", "code diff"),
    ("אלגוריתם ההפצה לאחור (backpropagation)", "backpropagation"),
    ("הפצה לאחור (backpropagation)", "backpropagation"),
    ("פונקציית אנטרופיה צולבת (cross\u00a0entropy)", "cross\u00a0entropy"),
    ("אנטרופיה צולבת (cross\u00a0entropy)", "cross\u00a0entropy"),
    ("פירוק לאסימונים (tokenization)", "tokenization"),
    ("מקטעים בני כמה אסימונים (n-grams)", "n-grams"),
    ("הטמעות (embeddings)", "embeddings"),
    ("מנגנון קשב (attention)", "attention"),
    ("הטמעת מיקום (positional\u00a0encoding)", "positional\u00a0encoding"),
    ("בנשירת יחידות (dropout)", "ב-dropout"),
    ("ונשירת יחידות (dropout)", "ו-dropout"),
    ("ובדעיכת משקלים (weight\u00a0decay)", "וב-weight\u00a0decay"),
    ("בתיוג חלש (weak labeling)", "בשיטת weak labeling"),
    ("מקטעי תווים קצרים (character\u00a0n-grams)", "character\u00a0n-grams"),
    ("הטמעה (embedding)", "embedding"),
    ("קשב עצמי (self-attention)", "self-attention"),
    ("דעיכת המשקלים (weight\u00a0decay)", "weight\u00a0decay"),
    ("דעיכת המשקלים", "weight\u00a0decay"),
    ("שיעור הדיוק (accuracy)", "accuracy"),
    ("דיוק ההתראות (precision)", "precision"),
    ("שיעור הכיסוי (recall)", "recall"),
    ("פירוק לאסימונים והטמעות", "tokenization ו-embeddings"),
    ("מנגנון קשב עם הטמעת מיקום", "attention עם positional\u00a0encoding"),
    ("מגבלת התיוג החלש", "מגבלת weak labeling"),
    ("וpositional", "ו-positional"),
    ("וcharacter", "ו-character"),
    ("לembedding", "ל-embedding"),
    ("וattention", "ו-attention"),
)


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def remove_all(parent, tag):
    for element in parent.findall(qn(tag)):
        parent.remove(element)


def set_rtl(r_pr):
    remove_all(r_pr, "w:rtl")
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    r_pr.append(rtl)


def set_english(r_pr):
    remove_all(r_pr, "w:rtl")
    remove_all(r_pr, "w:lang")
    lang = OxmlElement("w:lang")
    for attribute in ("w:val", "w:eastAsia", "w:bidi"):
        lang.set(qn(attribute), "en-US")
    r_pr.append(lang)


def replace_duplicate_terms(paragraph):
    text = paragraph.text
    for old, new in REPLACEMENTS:
        text = text.replace(old, new)
    if text == paragraph.text:
        return
    template = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    cursor = 0
    for match in ENGLISH_EXPRESSION.finditer(text):
        for value, english in ((text[cursor:match.start()], False), (match.group(), True)):
            if not value:
                continue
            run = paragraph.add_run(value)
            if template is not None:
                run._r.insert(0, deepcopy(template))
            r_pr = run._r.get_or_add_rPr()
            set_english(r_pr) if english else set_rtl(r_pr)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if template is not None:
            run._r.insert(0, deepcopy(template))
        set_rtl(run._r.get_or_add_rPr())


def main():
    path = Path(sys.argv[1] if len(sys.argv) > 1 else "report/DiffIntent_RTL_Final_Report_Hebrew.docx")
    document = Document(path)
    paragraphs = list(all_paragraphs(document))
    for paragraph in paragraphs:
        replace_duplicate_terms(paragraph)
        if paragraph.text == "5. תוצאות":
            paragraph.paragraph_format.page_break_before = True
    for index, paragraph in enumerate(paragraphs):
        has_hebrew = bool(HEBREW.search(paragraph.text.translate(DELETE_MARKS)))
        p_pr = paragraph._p.get_or_add_pPr()
        remove_all(p_pr, "w:bidi")
        if has_hebrew:
            p_pr.append(OxmlElement("w:bidi"))
            if index < 3 or paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.text = run.text.translate(DELETE_MARKS)
            r_pr = run._r.get_or_add_rPr()
            if LATIN.search(run.text) and not HEBREW.search(run.text):
                set_english(r_pr)
            elif has_hebrew and run.text:
                remove_all(r_pr, "w:lang")
                set_rtl(r_pr)
    temporary = path.with_suffix(".fixed.docx")
    document.save(temporary)
    temporary.replace(path)


if __name__ == "__main__":
    main()
